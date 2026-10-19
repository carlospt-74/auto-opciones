import json,re,glob,sys,os,datetime
sys.path.insert(0,'..'); sys.path.insert(0,os.path.dirname(os.path.abspath(__file__)))
from specs import parse
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
idx=json.load(open('index.json'))['data']['markdown']
index=re.findall(r'### (.+?)\n\n#### desde \$([\d,]+) MXN',idx)
parsed=json.load(open('parsed.json'))
BASE='https://nissanrivero.com/nissan/catalogo/nuevos/nissan-%s/'
PDF='https://d3s2hob8w3xwk8.cloudfront.net/autos-landing/nissan/%s/pdf/catalogo.pdf'
num=lambda s:int(re.sub(r'\D','',s))
def ts(slug):
    d=json.load(open(f'int_{slug}.json'));return d['ts']
wb=Workbook(); ws=wb.active; ws.title='Versiones'
H1=['Modelo','Año modelo','Versión','Precio MXN','Precio (texto en página)','Precio "desde" en listado del catálogo (modelo)','Precio "desde" en encabezado de la página (modelo)','Nota de precio','Equipamiento destacado (página web, textual)','Fuente del precio','URL página','URL PDF catálogo','Fecha/hora de captura (UTC)']
ws.append(H1)
ws2=wb.create_sheet('Fichas técnicas')
H2=['Modelo','Año modelo','Sección (PDF)','Concepto','Versión / columna (PDF)','Valor','Calidad de extracción','Fuente','Fecha de lectura (UTC)']
ws2.append(H2)
nver=nspec=0
for name,idxp in index:
    m=re.match(r'(.*) (\d{4})$',name); model,year=m.group(1),m.group(2)
    slug=name.lower().replace(' ','-')
    page=json.load(open(f'w2_{slug}.json'))['data']['markdown']
    hp=re.search(r'\ndesde\s+\$([\d,]+) MXN',page)
    head=num(hp.group(1)) if hp else None
    vers=parsed[slug]; t=ts(slug)
    prices=[num(v['precio']) for v in vers.values()]
    for v,x in vers.items():
        p=num(x['precio']); notes=[]
        if p==min(prices):
            if num(idxp)!=p: notes.append(f'El listado del catálogo dice desde ${num(idxp):,}; la versión más barata en la página es ${p:,}')
            if head and head!=p: notes.append(f'Encabezado de la página dice desde ${head:,}')
        if slug=='v-drive-2026': notes.append('El PDF de catálogo del sitio es de Nissan Colombia y no lista las versiones BASE TM / MT')
        if slug=='kait-2027': notes.append('El PDF incluye también la versión ADVANCE, que no aparece en la página')
        if not x.get('caracteristicas'): notes.append('La página no lista equipamiento para esta versión')
        if slug in('versa-2026','kait-2027','x-trail-2027'): notes.append('Precio confirmado por el usuario el 2026-10-19 (coincide con la página)')
        ws.append([model,int(year),v,p,x['precio'],num(idxp),head,'; '.join(notes),'\n'.join(x.get('caracteristicas',[])),'nissanrivero.com – página del modelo (pestaña de versión, vía Firecrawl interact)',BASE%slug,PDF%slug,t]); nver+=1
    # fichas
    tsr=datetime.datetime.fromtimestamp(os.path.getmtime(f'pdf_{slug}.json'),datetime.UTC).strftime('%Y-%m-%dT%H:%M:%SZ')
    if slug in('magnite-2026','x-trail-e-power-2026'):
        md=json.load(open(f'pdf_{slug}.json'))['data']['markdown']
        if slug=='magnite-2026': val='1.0 L, 5 cilindros turbo, 99 HP, 112 lb-ft (texto de marketing leído por OCR)'
        else: val='1.5 L turbo, 3 cilindros en línea, 142 HP (texto de marketing leído por OCR)'
        q='REVISAR: el PDF no trae tabla de especificaciones; solo texto de marketing con OCR deficiente'
        if slug=='magnite-2026': q+=' (el OCR dice "5 cilindros", dato dudoso)'
        ws2.append([model,int(year),'Motor','Motorización (texto del PDF)','Todas las versiones',val,q,PDF%slug,tsr]); nspec+=1
        continue
    rows,iss=parse(slug)
    for sec,con,ver,val,tag in rows:
        bad=(ver=='' or ver.startswith('col') or ver in('•','-','–','✓'))
        q='OK'
        if slug=='v-drive-2026': q='REVISAR: PDF de Nissan Colombia (solo V-DRIVE MT, cód. 67.101); puede no aplicar a las versiones mexicanas BASE TM / MT'
        elif bad: q='REVISAR: columna sin encabezado en la extracción'
        elif val=='': q='Celda vacía en la extracción (posible celda combinada del PDF)'
        elif tag=='iconlost': q='NO CONFIABLE: en el PDF esta tabla usa íconos y la extracción devolvió "-" en todas las versiones; no indica si el equipo está o no incluido'
        elif tag=='infer_all': q='Inferido: celda combinada; valor aplicado a todas las versiones de la tabla (confirmar visualmente en el PDF)'
        elif tag=='blank': q='REVISAR: fila con celdas combinadas; el valor puede aplicar a varias versiones'
        ws2.append([model,int(year),sec,con.replace('\\*','*'),ver,val,q,PDF%slug,tsr]); nspec+=1

for v,p in (('ADVANCE 2 ROW',529900),('EXCLUSIVE 2 ROW',591900),('PLATINUM 2 ROW',638900),('PLATINUM 3 ROW',661900),('PLATINUM PLUS 2 ROW',708900)):
    ws.append(['X-TRAIL',2026,v,p,f'Desde ${p:,} MXN',None,None,'Precio proporcionado por el usuario el 2026-10-19; no verificado en la página (solo se leyó la pestaña 2027). Sin equipamiento capturado.','','Proporcionado por el usuario (pestaña 2026 de la página del X-Trail, no leída por el scraping)',BASE%'x-trail-2027',PDF%'x-trail-2027','2026-10-19'])
    nver+=1

import epower2027
EP='fuentes/nissan-2027-x_trail_e_power-catalogo.pdf (PDF proporcionado por el usuario, X-Trail e-POWER 2027)'
for v,p,n in (('EXCLUSIVE e-POWER',751990,'El usuario indicó "Estrena desde $751,990" sin decir la versión; se asigna a la de entrada por ser precio "desde" (confirmar). El modelo 2027 es distinto de la fila 2026 (PLATINUM $842,900 en la página)'),('PLATINUM e-POWER e-4ORCE',None,'Precio no proporcionado')):
    ws.append(['X-TRAIL E-POWER',2027,v,p,f'Desde ${p:,} MXN' if p else None,None,None,n,'','Precio: dato del usuario ("Estrena desde"); versiones: PDF proporcionado por el usuario',None,EP,'2026-10-19']); nver+=1
for sec,con,ver,val,tag in epower2027.rows2():
    q='OK: celda combinada visible en el PDF; aplica a ambas versiones' if tag=='merged' else 'OK: texto del PDF con símbolos legibles (• aplica, - no aplica)'
    ws2.append(['X-TRAIL E-POWER',2027,sec,con,ver,val,q,EP,'2026-10-19']); nspec+=1
for w,H,widths in((ws,H1,[22,8,28,13,18,18,18,40,70,34,52,52,22]),(ws2,H2,[22,8,26,60,34,30,40,52,22])):
    for c in w[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='C3002F'); c.alignment=Alignment(wrap_text=True,vertical='center')
    for i,wd in enumerate(widths,1): w.column_dimensions[get_column_letter(i)].width=wd
    w.freeze_panes='A2'; w.auto_filter.ref=w.dimensions
for r in ws.iter_rows(min_row=2):
    r[3].number_format='"$"#,##0'; r[5].number_format='"$"#,##0'; r[6].number_format='"$"#,##0'
    for c in r: c.alignment=Alignment(wrap_text=True,vertical='top')
out='/home/user/auto-opciones/nissan-rivero/nissan_rivero_catalogo.xlsx'
wb.save(out); print(nver,nspec,out)
